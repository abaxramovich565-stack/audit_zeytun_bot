import asyncio
import logging
import re
from datetime import datetime, timedelta
from io import BytesIO

from aiogram import Bot, Dispatcher, Router, F
from aiogram.filters import Command, CommandStart, BaseFilter
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import (
    Message,
    CallbackQuery,
    InlineKeyboardMarkup,
    InlineKeyboardButton,
    BufferedInputFile,
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import config
import database as db

logging.basicConfig(level=logging.INFO)

bot = Bot(token=config.BOT_TOKEN)
dp = Dispatcher(storage=MemoryStorage())
router = Router()
dp.include_router(router)


# ---------------- FSM holatlari ----------------

class AuditFSM(StatesGroup):
    faol = State()               # Menyuda erkin harakatlanish
    izoh_kutilmoqda = State()    # Muammoga izoh/rasm kutilmoqda
    xodim_ism_kutilmoqda = State()  # Xodim ism-familiyasi kutilmoqda


def is_admin(user_id: int) -> bool:
    return user_id in config.ADMIN_IDS


# ---------------- Yordamchi funksiyalar ----------------

def filiallar_klaviaturasi():
    filiallar = db.filiallar_royxati()
    kb = InlineKeyboardMarkup(inline_keyboard=[])
    for f in filiallar:
        kb.inline_keyboard.append(
            [InlineKeyboardButton(text=f["nomi"], callback_data=f"filial:{f['id']}")]
        )
    return kb


def javob_klaviaturasi(band_id: int):
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(text="✅ Yaxshi", callback_data=f"javob:{band_id}:yaxshi"),
                InlineKeyboardButton(text="⚠️ O'rtacha", callback_data=f"javob:{band_id}:ortacha"),
                InlineKeyboardButton(text="❌ Yomon", callback_data=f"javob:{band_id}:yomon"),
            ],
            [InlineKeyboardButton(text="⬅️ Orqaga", callback_data="menyu")],
        ]
    )


def kategoriya_menyu_klaviaturasi(audit_id: int):
    katlar = db.kategoriyalar_royxati()
    javoblar = db.audit_javoblari_xaritasi(audit_id)
    kb = InlineKeyboardMarkup(inline_keyboard=[])
    for k in katlar:
        bandlar = db.bandlar_royxati(k["id"])
        jami = len(bandlar)
        javob_soni = sum(1 for b in bandlar if b["id"] in javoblar)
        belgi = "✅" if jami and javob_soni == jami else "🔸"
        matn = f"{belgi} {k['nomi']} ({javob_soni}/{jami})"
        kb.inline_keyboard.append([InlineKeyboardButton(text=matn[:60], callback_data=f"kat:{k['id']}")])
    kb.inline_keyboard.append(
        [InlineKeyboardButton(text="🏁 Auditni yakunlash", callback_data="yakunlash")]
    )
    return kb


def bandlar_menyu_klaviaturasi(kategoriya_id: int, audit_id: int):
    bandlar = db.bandlar_royxati(kategoriya_id)
    javoblar = db.audit_javoblari_xaritasi(audit_id)
    ikonka = {"yaxshi": "✅", "ortacha": "⚠️", "yomon": "❌"}
    kb = InlineKeyboardMarkup(inline_keyboard=[])
    for b in bandlar:
        belgi = ikonka.get(javoblar.get(b["id"]), "⬜️")
        matn = f"{belgi} {b['matn']}"
        kb.inline_keyboard.append([InlineKeyboardButton(text=matn[:60], callback_data=f"band:{b['id']}")])
    kb.inline_keyboard.append([InlineKeyboardButton(text="⬅️ Bosh menyu", callback_data="menyu")])
    return kb


def ball_hisoblash(holatlar):
    xarita = {"yaxshi": 100, "ortacha": 50, "yomon": 0}
    if not holatlar:
        return 0
    return round(sum(xarita[h] for h in holatlar) / len(holatlar), 1)


async def bosh_menyuni_korsatish(message: Message, audit_id: int):
    await message.answer(
        "Qaysi bo'limni tekshirmoqchisiz? Savollarni istalgan tartibda tanlashingiz mumkin.",
        reply_markup=kategoriya_menyu_klaviaturasi(audit_id),
    )


async def bandlar_royxatiga_qaytish(message: Message, kategoriya_id: int, audit_id: int):
    kat = db.kategoriya_olish(kategoriya_id)
    await message.answer(
        f"📂 {kat['nomi']}\n\nSavolni tanlang:",
        reply_markup=bandlar_menyu_klaviaturasi(kategoriya_id, audit_id),
    )


# ---------------- /start ----------------

@router.message(CommandStart())
async def cmd_start(message: Message, state: FSMContext):
    await state.clear()
    matn = (
        "Assalomu alaykum! Men supermarket auditi uchun botman.\n\n"
        "Buyruqlar:\n"
        "/audit — yangi tekshiruv boshlash\n"
        "/hisobot — oxirgi tekshiruvlar ro'yxati\n"
    )
    if is_admin(message.from_user.id):
        matn += (
            "\nAdmin buyruqlari:\n"
            "/filial_qoshish <nomi> — yangi filial qo'shish\n"
            "/filiallar — filiallar ro'yxati\n"
            "/kategoriyalar — checklist kategoriyalari\n"
            "/kategoriya_qoshish <nomi> — yangi kategoriya\n"
            "/band_qoshish <kategoriya_raqami> <savol> — yangi savol\n"
            "/bolim_qoshish <nomi> — xodimlar ishlaydigan bo'lim qo'shish\n"
            "/bolimlar — bo'limlar ro'yxati\n"
            "\nHisobotlar (Excel):\n"
            "/kunlik-hisobot — bugungi kun bo'yicha\n"
            "/haftalik-hisobot — oxirgi 7 kun bo'yicha\n"
            "/oylik-hisobot — oxirgi 30 kun bo'yicha\n"
        )
    await message.answer(matn)


# ---------------- Filial boshqaruvi (admin) ----------------

@router.message(Command("filial_qoshish"))
async def cmd_filial_qoshish(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Bu buyruq faqat adminlar uchun.")
        return
    nomi = message.text.replace("/filial_qoshish", "").strip()
    if not nomi:
        await message.answer("Filial nomini yozing. Masalan:\n/filial_qoshish Chilonzor filiali")
        return
    db.filial_qoshish(nomi)
    await message.answer(f"✅ '{nomi}' filiali qo'shildi.")


@router.message(Command("filiallar"))
async def cmd_filiallar(message: Message):
    filiallar = db.filiallar_royxati()
    if not filiallar:
        await message.answer("Hozircha filiallar qo'shilmagan. /filial_qoshish buyrug'idan foydalaning.")
        return
    matn = "📍 Filiallar:\n" + "\n".join(f"- {f['nomi']}" for f in filiallar)
    await message.answer(matn)


# ---------------- Bo'limlar boshqaruvi (admin) ----------------

@router.message(Command("bolim_qoshish"))
async def cmd_bolim_qoshish(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Bu buyruq faqat adminlar uchun.")
        return
    nomi = message.text.replace("/bolim_qoshish", "").strip()
    if not nomi:
        await message.answer("Format: /bolim_qoshish <nomi>\nMasalan: /bolim_qoshish Kassa")
        return
    db.bolim_qoshish(nomi)
    await message.answer(f"✅ '{nomi}' bo'limi qo'shildi.")


@router.message(Command("bolimlar"))
async def cmd_bolimlar(message: Message):
    bolimlar = db.bolimlar_royxati()
    if not bolimlar:
        await message.answer("Hozircha bo'limlar qo'shilmagan. /bolim_qoshish <nomi> orqali qo'shing.")
        return
    matn = "🏷 Bo'limlar:\n" + "\n".join(f"- {b['nomi']}" for b in bolimlar)
    await message.answer(matn)


# ---------------- Checklist boshqaruvi (admin) ----------------

@router.message(Command("kategoriyalar"))
async def cmd_kategoriyalar(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Bu buyruq faqat adminlar uchun.")
        return
    katlar = db.kategoriyalar_royxati()
    matn = "📂 Kategoriyalar:\n" + "\n".join(f"{k['id']}. {k['nomi']}" for k in katlar)
    matn += "\n\nYangi savol qo'shish uchun:\n/band_qoshish <kategoriya_raqami> <savol matni>"
    await message.answer(matn)


@router.message(Command("kategoriya_qoshish"))
async def cmd_kategoriya_qoshish(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Bu buyruq faqat adminlar uchun.")
        return
    nomi = message.text.replace("/kategoriya_qoshish", "").strip()
    if not nomi:
        await message.answer("Format: /kategoriya_qoshish <nomi>")
        return
    db.kategoriya_qoshish(nomi)
    await message.answer(f"✅ '{nomi}' kategoriyasi qo'shildi.")


@router.message(Command("band_qoshish"))
async def cmd_band_qoshish(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Bu buyruq faqat adminlar uchun.")
        return
    qism = message.text.split(maxsplit=2)
    if len(qism) < 3 or not qism[1].isdigit():
        await message.answer(
            "To'g'ri format:\n/band_qoshish <kategoriya_raqami> <savol matni>\n"
            "Kategoriyalar ro'yxatini /kategoriyalar orqali ko'ring."
        )
        return
    kategoriya_id, matn = int(qism[1]), qism[2]
    db.band_qoshish(kategoriya_id, matn)
    await message.answer("✅ Yangi savol qo'shildi.")


# ---------------- Audit jarayoni (erkin tartib) ----------------

@router.message(Command("audit"))
async def cmd_audit(message: Message, state: FSMContext):
    filiallar = db.filiallar_royxati()
    if not filiallar:
        await message.answer(
            "Hali birorta filial qo'shilmagan.\n"
            "Admin /filial_qoshish <nomi> buyrug'i orqali filial qo'shishi kerak."
        )
        return
    await state.clear()
    await message.answer("Qaysi filialni tekshiryapsiz?", reply_markup=filiallar_klaviaturasi())


@router.callback_query(F.data.startswith("filial:"))
async def filial_tanlandi(callback: CallbackQuery, state: FSMContext):
    filial_id = int(callback.data.split(":")[1])
    filial = db.filial_olish(filial_id)

    kategoriyalar = db.kategoriyalar_royxati()
    if not kategoriyalar:
        await callback.message.answer("Checklist kategoriyalari topilmadi.")
        await callback.answer()
        return

    audit_id = db.audit_boshlash(
        filial_id=filial_id,
        auditor_id=callback.from_user.id,
        auditor_ism=callback.from_user.full_name,
        vaqt=datetime.now().isoformat(timespec="seconds"),
    )

    await state.set_state(AuditFSM.faol)
    await state.update_data(audit_id=audit_id, filial_id=filial_id)

    await callback.message.answer(f"🏬 Filial: {filial['nomi']}\nTekshiruv boshlandi!")
    await bosh_menyuni_korsatish(callback.message, audit_id)
    await callback.answer()


@router.callback_query(AuditFSM.faol, F.data == "menyu")
async def bosh_menyuga_qaytish(callback: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    await callback.message.edit_text(
        "Qaysi bo'limni tekshirmoqchisiz? Savollarni istalgan tartibda tanlashingiz mumkin.",
        reply_markup=kategoriya_menyu_klaviaturasi(data["audit_id"]),
    )
    await callback.answer()


@router.callback_query(AuditFSM.faol, F.data.startswith("kat:"))
async def kategoriya_tanlandi(callback: CallbackQuery, state: FSMContext):
    kategoriya_id = int(callback.data.split(":")[1])
    data = await state.get_data()
    kat = db.kategoriya_olish(kategoriya_id)
    await callback.message.edit_text(
        f"📂 {kat['nomi']}\n\nSavolni tanlang:",
        reply_markup=bandlar_menyu_klaviaturasi(kategoriya_id, data["audit_id"]),
    )
    await callback.answer()


@router.callback_query(AuditFSM.faol, F.data.startswith("band:"))
async def band_tanlandi(callback: CallbackQuery, state: FSMContext):
    band_id = int(callback.data.split(":")[1])
    band = db.band_kategoriya_olish(band_id)
    await callback.message.edit_text(
        f"[{band['kategoriya_nomi']}]\n\n{band['matn']}",
        reply_markup=javob_klaviaturasi(band_id),
    )
    await callback.answer()


@router.callback_query(AuditFSM.faol, F.data.startswith("javob:"))
async def javob_qabul_qilindi(callback: CallbackQuery, state: FSMContext):
    _, band_id, holat = callback.data.split(":")
    band_id = int(band_id)
    band = db.band_kategoriya_olish(band_id)
    data = await state.get_data()
    audit_id = data["audit_id"]

    # Xodimlar ishi bo'limi: rasm/izoh talab qilinmaydi, buning o'rniga
    # xodim qaysi bo'limda ishlashi va ism-familiyasi so'raladi.
    if band["xodim_bolimi"]:
        bolimlar = db.bolimlar_royxati()
        if not bolimlar:
            db.javob_saqlash(audit_id=audit_id, band_id=band_id, holat=holat)
            await callback.message.edit_text(
                "✅ Saqlandi.\n\n"
                "ℹ️ Bo'limlar ro'yxati hali bo'sh. Admin /bolim_qoshish <nomi> "
                "orqali bo'lim qo'shsa, keyingi safar xodim ishlaydigan bo'lim "
                "va ismi so'raladi."
            )
            await bandlar_royxatiga_qaytish(callback.message, band["kategoriya_id"], audit_id)
            await callback.answer()
            return

        await state.update_data(kutilayotgan_band=band_id, kutilayotgan_holat=holat)
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text=b["nomi"], callback_data=f"bolim:{b['id']}")]
            for b in bolimlar
        ])
        kb.inline_keyboard.append([InlineKeyboardButton(text="⬅️ Orqaga", callback_data="menyu")])
        await callback.message.edit_text("Xodim qaysi bo'limda ishlaydi?", reply_markup=kb)
        await callback.answer()
        return

    if holat == "yaxshi":
        db.javob_saqlash(audit_id=audit_id, band_id=band_id, holat=holat)
        await callback.message.edit_text("✅ Saqlandi.")
        await bandlar_royxatiga_qaytish(callback.message, band["kategoriya_id"], audit_id)
    else:
        await state.update_data(kutilayotgan_band=band_id, kutilayotgan_holat=holat)
        await callback.message.edit_text(
            "📸 Muammo haqida izoh yozing yoki rasm yuboring.\n"
            "Izohsiz saqlamoqchi bo'lsangiz — /otkazib_yuborish deb yozing."
        )
        await state.set_state(AuditFSM.izoh_kutilmoqda)

    await callback.answer()


@router.callback_query(AuditFSM.faol, F.data.startswith("bolim:"))
async def bolim_tanlandi(callback: CallbackQuery, state: FSMContext):
    bolim_id = int(callback.data.split(":")[1])
    await state.update_data(kutilayotgan_bolim=bolim_id)
    await state.set_state(AuditFSM.xodim_ism_kutilmoqda)
    await callback.message.edit_text("Xodimning ism-familiyasini yozing:")
    await callback.answer()


@router.message(AuditFSM.xodim_ism_kutilmoqda, F.text)
async def xodim_ism_qabul_qilindi(message: Message, state: FSMContext):
    data = await state.get_data()
    ism = message.text.strip()
    db.javob_saqlash(
        audit_id=data["audit_id"],
        band_id=data["kutilayotgan_band"],
        holat=data["kutilayotgan_holat"],
        bolim_id=data["kutilayotgan_bolim"],
        xodim_ism=ism,
    )
    band = db.band_kategoriya_olish(data["kutilayotgan_band"])
    await message.answer(f"✅ Saqlandi: {ism}")
    await state.set_state(AuditFSM.faol)
    await bandlar_royxatiga_qaytish(message, band["kategoriya_id"], data["audit_id"])


@router.message(AuditFSM.izoh_kutilmoqda, Command("otkazib_yuborish"))
async def izohsiz_otish(message: Message, state: FSMContext):
    await izohni_saqlash_va_qaytish(message, state, izoh=None, rasm_file_id=None)


@router.message(AuditFSM.izoh_kutilmoqda, F.photo)
async def rasm_bilan_izoh(message: Message, state: FSMContext):
    rasm_file_id = message.photo[-1].file_id
    izoh = message.caption
    await izohni_saqlash_va_qaytish(message, state, izoh=izoh, rasm_file_id=rasm_file_id)


@router.message(AuditFSM.izoh_kutilmoqda, F.text)
async def matn_bilan_izoh(message: Message, state: FSMContext):
    await izohni_saqlash_va_qaytish(message, state, izoh=message.text, rasm_file_id=None)


async def izohni_saqlash_va_qaytish(message: Message, state: FSMContext, izoh, rasm_file_id):
    data = await state.get_data()
    db.javob_saqlash(
        audit_id=data["audit_id"],
        band_id=data["kutilayotgan_band"],
        holat=data["kutilayotgan_holat"],
        izoh=izoh,
        rasm_file_id=rasm_file_id,
    )
    band = db.band_kategoriya_olish(data["kutilayotgan_band"])
    await message.answer("✅ Saqlandi.")
    await state.set_state(AuditFSM.faol)
    await bandlar_royxatiga_qaytish(message, band["kategoriya_id"], data["audit_id"])


@router.callback_query(AuditFSM.faol, F.data == "yakunlash")
async def yakunlashni_sorash(callback: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    audit_id = data["audit_id"]
    jami = len(db.barcha_bandlarni_yigish())
    javob_soni = len(db.audit_holatlar_royxati(audit_id))

    if javob_soni < jami:
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="✅ Ha, yakunlash", callback_data="yakunlash_ha")],
            [InlineKeyboardButton(text="⬅️ Yo'q, davom etish", callback_data="menyu")],
        ])
        await callback.message.edit_text(
            f"⚠️ Hali {jami - javob_soni} ta savol javobsiz qoldi.\n"
            "Baribir yakunlaysizmi?",
            reply_markup=kb,
        )
    else:
        await auditni_yakunlash(callback.message, state)
    await callback.answer()


@router.callback_query(AuditFSM.faol, F.data == "yakunlash_ha")
async def yakunlashni_tasdiqlash(callback: CallbackQuery, state: FSMContext):
    await auditni_yakunlash(callback.message, state)
    await callback.answer()


async def auditni_yakunlash(message: Message, state: FSMContext):
    data = await state.get_data()
    holatlar = db.audit_holatlar_royxati(data["audit_id"])
    umumiy_ball = ball_hisoblash(holatlar)
    db.audit_yakunlash(
        audit_id=data["audit_id"],
        vaqt=datetime.now().isoformat(timespec="seconds"),
        umumiy_ball=umumiy_ball,
    )
    muammolar_soni = sum(1 for h in holatlar if h != "yaxshi")
    await message.answer(
        f"🏁 Tekshiruv yakunlandi!\n\n"
        f"📊 Umumiy ball: {umumiy_ball}%\n"
        f"⚠️ Aniqlangan muammolar: {muammolar_soni} ta\n\n"
        f"Rahmat! /hisobot orqali batafsil ko'rishingiz mumkin."
    )
    await state.clear()


# ---------------- Hisobot (bitta audit bo'yicha, Telegram xabarida) ----------------

@router.message(Command("hisobot"))
async def cmd_hisobot(message: Message):
    auditlar = db.oxirgi_auditlar(limit=10)
    if not auditlar:
        await message.answer("Hali yakunlangan tekshiruvlar yo'q.")
        return

    kb = InlineKeyboardMarkup(inline_keyboard=[])
    matn = "📋 Oxirgi tekshiruvlar:\n\n"
    for a in auditlar:
        sana = a["tugash_vaqti"][:16].replace("T", " ")
        matn += f"🏬 {a['filial_nomi']} — {sana} — {a['umumiy_ball']}%\n"
        kb.inline_keyboard.append(
            [InlineKeyboardButton(
                text=f"{a['filial_nomi']} ({sana})",
                callback_data=f"hisobot:{a['id']}",
            )]
        )
    await message.answer(matn, reply_markup=kb)


@router.callback_query(F.data.startswith("hisobot:"))
async def hisobot_batafsil(callback: CallbackQuery):
    audit_id = int(callback.data.split(":")[1])
    audit, javoblar = db.audit_toliq(audit_id)

    matn = (
        f"🏬 Filial: {audit['filial_nomi']}\n"
        f"👤 Auditor: {audit['auditor_ism']}\n"
        f"📅 Sana: {audit['tugash_vaqti'][:16].replace('T', ' ')}\n"
        f"📊 Umumiy ball: {audit['umumiy_ball']}%\n\n"
        f"Muammolar:"
    )
    await callback.message.answer(matn)

    muammolar = [j for j in javoblar if j["holat"] != "yaxshi"]
    if not muammolar:
        await callback.message.answer("Muammo aniqlanmadi. 🎉")
    else:
        for j in muammolar:
            belgi = "❌" if j["holat"] == "yomon" else "⚠️"
            qoshimcha = ""
            if j["xodim_ism"]:
                qoshimcha += f"\n👤 Xodim: {j['xodim_ism']}"
            if j["bolim_nomi"]:
                qoshimcha += f"\n🏷 Bo'lim: {j['bolim_nomi']}"
            izoh_matni = f"\nIzoh: {j['izoh']}" if j["izoh"] else ""
            javob_matni = f"{belgi} {j['band_matni']}{qoshimcha}{izoh_matni}"
            if j["rasm_file_id"]:
                await callback.message.answer_photo(j["rasm_file_id"], caption=javob_matni)
            else:
                await callback.message.answer(javob_matni)

    await callback.answer()


# ---------------- Davriy hisobotlar (Excel) ----------------

_SARLAVHA_FONT = Font(bold=True, color="FFFFFF", size=11)
_SARLAVHA_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
_YAXSHI_RANG = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_ORTACHA_RANG = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
_YOMON_RANG = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


def _chegara():
    yon = Side(style="thin", color="B7B7B7")
    return Border(left=yon, right=yon, top=yon, bottom=yon)


def _ball_rangi(ball):
    if ball >= 80:
        return _YAXSHI_RANG
    if ball >= 50:
        return _ORTACHA_RANG
    return _YOMON_RANG


def _ustunlar_kengligini_moslashtirish(ws):
    for col in ws.columns:
        uzunlik = 0
        harf = None
        for cell in col:
            if harf is None:
                harf = get_column_letter(cell.column)
            qiymat = cell.value
            if qiymat is not None:
                uzunlik = max(uzunlik, len(str(qiymat)))
        if harf:
            ws.column_dimensions[harf].width = min(max(uzunlik + 2, 10), 55)


def _sarlavha_qatorini_yozish(ws, ustunlar):
    for i, nomi in enumerate(ustunlar, start=1):
        c = ws.cell(row=1, column=i, value=nomi)
        c.font = _SARLAVHA_FONT
        c.fill = _SARLAVHA_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _chegara()
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24


def hisobot_excel_yaratish(boshlanish_iso: str, tugash_iso: str) -> bytes:
    auditlar = db.auditlar_oraliqda(boshlanish_iso, tugash_iso)

    wb = Workbook()
    xulosa_ws = wb.active
    xulosa_ws.title = "Xulosa"
    _sarlavha_qatorini_yozish(
        xulosa_ws, ["Filial", "Tekshiruvlar soni", "O'rtacha ball (%)", "Aniqlangan muammolar"]
    )

    filial_stat = {}
    javoblar_kesh = {}
    for a in auditlar:
        _, javoblar = db.audit_toliq(a["id"])
        javoblar_kesh[a["id"]] = javoblar

        s = filial_stat.setdefault(a["filial_nomi"], {"soni": 0, "ball_yigindisi": 0.0, "muammo": 0})
        s["soni"] += 1
        s["ball_yigindisi"] += a["umumiy_ball"] or 0
        s["muammo"] += sum(1 for j in javoblar if j["holat"] != "yaxshi")

    qator = 2
    for nomi, s in sorted(filial_stat.items()):
        ortacha = round(s["ball_yigindisi"] / s["soni"], 1) if s["soni"] else 0
        xulosa_ws.cell(row=qator, column=1, value=nomi).border = _chegara()
        xulosa_ws.cell(row=qator, column=2, value=s["soni"]).border = _chegara()
        ball_katak = xulosa_ws.cell(row=qator, column=3, value=ortacha)
        ball_katak.border = _chegara()
        ball_katak.fill = _ball_rangi(ortacha)
        ball_katak.alignment = Alignment(horizontal="center")
        xulosa_ws.cell(row=qator, column=4, value=s["muammo"]).border = _chegara()
        qator += 1

    if not filial_stat:
        xulosa_ws.cell(row=2, column=1, value="Bu davrda yakunlangan tekshiruv topilmadi.")

    _ustunlar_kengligini_moslashtirish(xulosa_ws)

    tafsilot_ws = wb.create_sheet("Tafsilot")
    _sarlavha_qatorini_yozish(
        tafsilot_ws,
        ["Filial", "Sana", "Auditor", "Umumiy ball (%)", "Kategoriya",
         "Savol", "Holat", "Bo'lim", "Xodim", "Izoh"],
    )

    holat_matni = {"yaxshi": "✅ Yaxshi", "ortacha": "⚠️ O'rtacha", "yomon": "❌ Yomon"}
    holat_rangi = {"ortacha": _ORTACHA_RANG, "yomon": _YOMON_RANG}

    qator = 2
    for a in auditlar:
        sana = (a["tugash_vaqti"] or "")[:16].replace("T", " ")
        for j in javoblar_kesh[a["id"]]:
            tafsilot_ws.cell(row=qator, column=1, value=a["filial_nomi"]).border = _chegara()
            tafsilot_ws.cell(row=qator, column=2, value=sana).border = _chegara()
            tafsilot_ws.cell(row=qator, column=3, value=a["auditor_ism"] or "").border = _chegara()
            tafsilot_ws.cell(row=qator, column=4, value=a["umumiy_ball"]).border = _chegara()
            tafsilot_ws.cell(row=qator, column=5, value=j["kategoriya_nomi"]).border = _chegara()
            tafsilot_ws.cell(row=qator, column=6, value=j["band_matni"]).border = _chegara()
            holat_katak = tafsilot_ws.cell(row=qator, column=7, value=holat_matni.get(j["holat"], j["holat"]))
            holat_katak.border = _chegara()
            if j["holat"] in holat_rangi:
                holat_katak.fill = holat_rangi[j["holat"]]
            tafsilot_ws.cell(row=qator, column=8, value=j["bolim_nomi"] or "").border = _chegara()
            tafsilot_ws.cell(row=qator, column=9, value=j["xodim_ism"] or "").border = _chegara()
            tafsilot_ws.cell(row=qator, column=10, value=j["izoh"] or "").border = _chegara()
            qator += 1

    if qator == 2:
        tafsilot_ws.cell(row=2, column=1, value="Bu davrda ma'lumot topilmadi.")

    _ustunlar_kengligini_moslashtirish(tafsilot_ws)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class DavriyHisobotFiltri(BaseFilter):
    """/kunlik-hisobot, /haftalik-hisobot, /oylik-hisobot (yoki _ bilan) ni ushlaydi."""

    NAQSH = re.compile(r"^/(kunlik|haftalik|oylik)[-_]hisobot(?:@\w+)?$")

    async def __call__(self, message: Message) -> bool:
        return bool(message.text and self.NAQSH.match(message.text))


@router.message(DavriyHisobotFiltri())
async def cmd_davriy_hisobot(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Bu buyruq faqat adminlar uchun.")
        return

    turi = DavriyHisobotFiltri.NAQSH.match(message.text).group(1)
    hozir = datetime.now()

    if turi == "kunlik":
        boshlanish = hozir.replace(hour=0, minute=0, second=0, microsecond=0)
        sarlavha = "Kunlik hisobot"
    elif turi == "haftalik":
        boshlanish = hozir - timedelta(days=7)
        sarlavha = "Haftalik hisobot"
    else:
        boshlanish = hozir - timedelta(days=30)
        sarlavha = "Oylik hisobot"

    kutish_xabari = await message.answer("⏳ Hisobot tayyorlanmoqda...")

    excel_bytes = hisobot_excel_yaratish(
        boshlanish.isoformat(timespec="seconds"),
        hozir.isoformat(timespec="seconds"),
    )
    fayl_nomi = f"{turi}_hisobot_{hozir.strftime('%Y-%m-%d')}.xlsx"
    hujjat = BufferedInputFile(excel_bytes, filename=fayl_nomi)

    await message.answer_document(
        hujjat,
        caption=(
            f"📊 {sarlavha}\n"
            f"{boshlanish.strftime('%d.%m.%Y')} — {hozir.strftime('%d.%m.%Y')}"
        ),
    )
    await kutish_xabari.delete()


# ---------------- Ishga tushirish ----------------

async def main():
    db.init_db()
    print("Bot ishga tushdi...")
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
